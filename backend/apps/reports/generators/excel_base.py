"""
apps/reports/generators/excel_base.py
Base class for all Excel report generators using openpyxl.
"""

import io
from django.utils import timezone


class BaseExcelGenerator:

    TITLE      = "CallSight AI Report"
    SHEET_NAME = "Report"

    HEADER_BG    = "4A4A9C"
    HEADER_FG    = "FFFFFF"
    SUBHEADER_BG = "E8E8F5"
    SUBHEADER_FG = "4A4A9C"
    ACCENT_BG    = "F5F5FF"
    BORDER_COLOR = "CCCCCC"
    POSITIVE_BG  = "D4EDDA"
    NEGATIVE_BG  = "F8D7DA"
    NEUTRAL_BG   = "E2E3E5"

    def __init__(self, company, date_from, date_to):
        self.company   = company
        self.date_from = date_from
        self.date_to   = date_to

    def _imports(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                PatternFill, Font, Alignment, Border, Side, GradientFill
            )
            from openpyxl.utils import get_column_letter
            return Workbook, PatternFill, Font, Alignment, Border, Side, get_column_letter
        except ImportError:
            raise ImportError("openpyxl is not installed. Run: pip install openpyxl")

    # ── Style helpers ──────────────────────────────────────────────────────────

    def _fill(self, PatternFill, color):
        return PatternFill("solid", fgColor=color)

    def _font(self, Font, bold=False, color="000000", size=10, italic=False):
        return Font(bold=bold, color=color, size=size, italic=italic)

    def _align(self, Alignment, horizontal="left", vertical="center", wrap=False):
        return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)

    def _border(self, Border, Side):
        thin = Side(style="thin", color=self.BORDER_COLOR)
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Row writers ───────────────────────────────────────────────────────────

    def _write_title(self, ws, Font, Alignment):
        ws.append([f"CallSight AI — {self.TITLE}"])
        cell       = ws.cell(row=ws.max_row, column=1)
        cell.font  = Font(bold=True, size=16, color=self.HEADER_BG)
        cell.alignment = Alignment(horizontal="left")
        ws.append([])  # blank row

    def _write_meta(self, ws, Font, Alignment):
        ws.append([
            f"Company: {self.company.name}",
            f"Period: {self.date_from} → {self.date_to}",
            f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M UTC')}",
        ])
        row = ws.max_row
        for col in range(1, 4):
            cell           = ws.cell(row=row, column=col)
            cell.font      = Font(italic=True, color="888888", size=9)
            cell.alignment = Alignment(horizontal="left")
        ws.append([])  # blank row

    def _write_section_title(self, ws, title, Font, Alignment):
        ws.append([title])
        cell           = ws.cell(row=ws.max_row, column=1)
        cell.font      = Font(bold=True, size=12, color=self.HEADER_BG)
        cell.alignment = Alignment(horizontal="left")

    def _write_headers(self, ws, headers, PatternFill, Font, Alignment, Border, Side):
        ws.append(headers)
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell           = ws.cell(row=row, column=col)
            cell.fill      = self._fill(PatternFill, self.HEADER_BG)
            cell.font      = self._font(Font, bold=True, color=self.HEADER_FG, size=10)
            cell.alignment = self._align(Alignment, horizontal="center")
            cell.border    = self._border(Border, Side)

    def _write_row(self, ws, data, PatternFill, Font, Alignment, Border, Side,
                   bg_color=None):
        ws.append(data)
        row = ws.max_row
        for col in range(1, len(data) + 1):
            cell           = ws.cell(row=row, column=col)
            cell.alignment = self._align(Alignment)
            cell.border    = self._border(Border, Side)
            if bg_color:
                cell.fill  = self._fill(PatternFill, bg_color)

    def _write_kpi_section(self, ws, kpis, PatternFill, Font, Alignment, Border, Side):
        """
        kpis = list of (label, value) tuples
        Writes two rows: labels + values, styled.
        """
        labels = [k[0] for k in kpis]
        values = [k[1] for k in kpis]

        ws.append(labels)
        row = ws.max_row
        for col in range(1, len(labels) + 1):
            cell           = ws.cell(row=row, column=col)
            cell.fill      = self._fill(PatternFill, self.SUBHEADER_BG)
            cell.font      = self._font(Font, bold=True, color=self.SUBHEADER_FG, size=9)
            cell.alignment = self._align(Alignment, horizontal="center")
            cell.border    = self._border(Border, Side)

        ws.append(values)
        row = ws.max_row
        for col in range(1, len(values) + 1):
            cell           = ws.cell(row=row, column=col)
            cell.fill      = self._fill(PatternFill, self.ACCENT_BG)
            cell.font      = self._font(Font, bold=True, color=self.HEADER_BG, size=14)
            cell.alignment = self._align(Alignment, horizontal="center")
            cell.border    = self._border(Border, Side)

        ws.append([])  # blank row

    def _auto_width(self, ws):
        for col in ws.columns:
            max_len    = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    def _freeze_pane(self, ws, cell="A2"):
        ws.freeze_panes = cell

    def _to_bytes(self, wb):
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def generate(self):
        raise NotImplementedError("Subclasses must implement generate()")